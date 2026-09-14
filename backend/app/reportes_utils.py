"""
Utilidades compartidas para generar reportes en PDF (reportlab) y
Excel (openpyxl) a partir de datos tabulares (encabezados + filas),
con una presentación corporativa: logo de la empresa, nombre, fecha y
hora de generación, usuario que lo generó y periodo consultado.
"""

import os
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

EMPRESA_NOMBRE = "JHM Tech Solutions"
EMPRESA_NIT = "NIT 900.000.000-1"
COLOR_PRIMARIO = colors.HexColor("#1d4ed8")
RUTA_LOGO = os.path.join(os.path.dirname(__file__), "static", "logo.png")


def _formatear_periodo(fecha_desde=None, fecha_hasta=None) -> str:
    if fecha_desde and fecha_hasta:
        return f"Del {fecha_desde.strftime('%d/%m/%Y')} al {fecha_hasta.strftime('%d/%m/%Y')}"
    if fecha_desde:
        return f"Desde el {fecha_desde.strftime('%d/%m/%Y')}"
    if fecha_hasta:
        return f"Hasta el {fecha_hasta.strftime('%d/%m/%Y')}"
    return "Todos los registros"


def generar_pdf_reporte(
    titulo: str,
    headers: list[str],
    filas: list[list],
    generado_por: str = "Sistema",
    fecha_desde=None,
    fecha_hasta=None,
    totales: list[list] | None = None,
) -> BytesIO:
    """Genera un PDF tabular corporativo (logo + metadatos + tabla) en memoria."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        topMargin=1.3 * cm,
        bottomMargin=1.3 * cm,
        leftMargin=1.3 * cm,
        rightMargin=1.3 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = []

    # ---- Encabezado con logo ----
    if os.path.exists(RUTA_LOGO):
        encabezado = Table(
            [[Image(RUTA_LOGO, width=2.6 * cm, height=2.2 * cm), Paragraph(
                f"<b>{EMPRESA_NOMBRE}</b><br/>{EMPRESA_NIT}<br/>{titulo}",
                estilos["Normal"],
            )]],
            colWidths=[3 * cm, 20 * cm],
        )
        encabezado.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        elementos.append(encabezado)
    else:
        elementos.append(Paragraph(EMPRESA_NOMBRE, estilos["Heading1"]))
        elementos.append(Paragraph(titulo, estilos["Heading2"]))

    elementos.append(Spacer(1, 0.4 * cm))
    metadatos = (
        f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} &nbsp;|&nbsp; "
        f"Generado por: {generado_por} &nbsp;|&nbsp; "
        f"Periodo: {_formatear_periodo(fecha_desde, fecha_hasta)}"
    )
    elementos.append(Paragraph(metadatos, estilos["Italic"]))
    elementos.append(Spacer(1, 0.5 * cm))

    data = [headers] + [[str(celda) if celda is not None else "" for celda in fila] for fila in filas]
    tabla = Table(data, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), COLOR_PRIMARIO),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    elementos.append(tabla)

    if not filas:
        elementos.append(Spacer(1, 0.5 * cm))
        elementos.append(Paragraph("No se encontraron registros para los filtros aplicados.", estilos["Normal"]))

    if totales:
        elementos.append(Spacer(1, 0.5 * cm))
        tabla_totales = Table(totales, colWidths=[6 * cm, 4 * cm])
        tabla_totales.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("LINEABOVE", (0, 0), (-1, 0), 1, colors.black),
                ]
            )
        )
        elementos.append(tabla_totales)

    doc.build(elementos)
    buffer.seek(0)
    return buffer


def generar_excel_reporte(
    titulo: str,
    headers: list[str],
    filas: list[list],
    generado_por: str = "Sistema",
    fecha_desde=None,
    fecha_hasta=None,
) -> BytesIO:
    """Genera un archivo Excel (.xlsx) corporativo (logo + metadatos + tabla)."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] or "Reporte"

    num_columnas = max(len(headers), 1)
    fila_actual = 1

    if os.path.exists(RUTA_LOGO):
        img = ExcelImage(RUTA_LOGO)
        img.height = 60
        img.width = 70
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 46
        ws.row_dimensions[2].height = 18
        fila_actual = 4

    ws.cell(row=fila_actual, column=1, value=EMPRESA_NOMBRE).font = Font(bold=True, size=13)
    fila_actual += 1
    ws.cell(row=fila_actual, column=1, value=titulo).font = Font(bold=True, size=11, color="1D4ED8")
    fila_actual += 1
    metadatos = (
        f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} | "
        f"Por: {generado_por} | Periodo: {_formatear_periodo(fecha_desde, fecha_hasta)}"
    )
    ws.cell(row=fila_actual, column=1, value=metadatos).font = Font(italic=True, size=9, color="6B7280")
    fila_actual += 2

    fila_headers = fila_actual
    for col_idx, header in enumerate(headers, start=1):
        celda = ws.cell(row=fila_headers, column=col_idx, value=header)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        fila_actual += 1
        for col_idx, valor in enumerate(fila, start=1):
            ws.cell(row=fila_actual, column=col_idx, value=valor)

    for col_idx, header in enumerate(headers, start=1):
        letra = get_column_letter(col_idx)
        ancho = max(len(str(header)) + 4, 14)
        ws.column_dimensions[letra].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
